# -*- coding: utf-8 -*-
"""
인천광역시교육청 「학교 현황」 엑셀(고등학교 시트)로 학교 목록·구·군을 최신화합니다.

사용법 (저장소 루트에서):
    pip install openpyxl
    python tools/sync_school_list.py "2026_4_1_자_학교_현황(최종).xlsx"
    python tools/sync_school_list.py 현황.xlsx --district-col 구군별      # 개편 전 구·군으로
    python tools/sync_school_list.py 현황.xlsx --include 일반,자율형,특목,특성화

하는 일:
  1. tools/districts.csv 를 공식 현황 기준으로 다시 만듭니다 (전체 고등학교).
  2. schools.json 을 다시 정리합니다.
     - 모든 학교를 공식 구·군으로 옮김 (잘못 분류된 학교 바로잡기)
     - 편성표 엑셀(schools/*.xlsx)이 있는 학교는 그대로 선택 가능
     - 엑셀이 없는 학교도 목록에 넣고 "pending": true (사이트에 '편성표 준비 중'으로 표시)
     - 기존 학교 id·파일 경로·수정일은 그대로 유지
  3. tools/missing_curriculum.csv 에 편성표가 아직 없는 학교를 적어 둡니다.
     → python tools/fetch_ice_curriculum.py --only 작전여자,도림 처럼 받아 오면 자동으로 선택 가능해집니다.

구·군 기준:
  기본값은 '개편후 행정구역(2026.7.1.시행)' 열입니다 (제물포구·영종구·서해구·검단구).
  개편 전 이름(중구·동구·서구)을 쓰려면 --district-col 구군별
"""
import argparse, csv, json, os, re, sys, time

try:
    import openpyxl
except ImportError:
    sys.exit("먼저 실행:  pip install openpyxl")

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from fetch_ice_curriculum import make_id, short_name   # 같은 id 규칙 사용

HEADER_KEYS = {"계열별": "cat", "구군별": "gu", "개편후": "newgu", "설립별": "est", "남/여/공학": "sex", "학교명": "name"}
# 공식 현황에는 없지만 편성표가 따로 올라오는 학과 단위 항목 → 본교와 같은 구·군
ALIASES = {"인천외국어고등학교(국제계열)": "인천외국어고등학교"}


def read_official(path, sheet="고등학교"):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        sys.exit(f"'{sheet}' 시트를 찾지 못했어요. 시트 목록: {wb.sheetnames}")
    ws = wb[sheet]
    col, rows = {}, []
    for r in ws.iter_rows(values_only=True):
        if not col:
            cells = [str(c or "").replace("\n", "") for c in r]
            if "학교명" in cells:
                for i, c in enumerate(cells):
                    for k, v in HEADER_KEYS.items():
                        if c.startswith(k) and v not in col:
                            col[v] = i
            continue
        name = r[col["name"]] if len(r) > col["name"] else None
        if not isinstance(name, str) or not name.strip().endswith("학교"):
            continue                                   # 소계·합계 행 건너뜀
        g = lambda k: (str(r[col[k]]).strip() if k in col and r[col[k]] is not None else "")
        rows.append({"name": name.strip(), "cat": g("cat"), "gu": g("gu"), "newgu": g("newgu") or g("gu"),
                     "est": g("est"), "sex": g("sex")})
    if not rows:
        sys.exit("학교 행을 읽지 못했어요. 시트 머리글(학교명·구군별 등)을 확인해 주세요.")
    return rows


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx", help="교육청 학교 현황 엑셀 경로")
    ap.add_argument("--district-col", choices=["개편후", "구군별"], default="개편후")
    ap.add_argument("--include", default="일반,자율형,특목", help="사이트 학교 목록에 넣을 계열 (쉼표)")
    ap.add_argument("--json", default="schools.json")
    ap.add_argument("--schools-dir", default="schools")
    ap.add_argument("--districts", default="tools/districts.csv")
    ap.add_argument("--missing", default="tools/missing_curriculum.csv")
    a = ap.parse_args()

    official = read_official(a.xlsx)
    key = "newgu" if a.district_col == "개편후" else "gu"
    by_name = {o["name"]: o for o in official}
    include = {x.strip() for x in a.include.split(",") if x.strip()}

    # 1) districts.csv
    os.makedirs(os.path.dirname(a.districts) or ".", exist_ok=True)
    with open(a.districts, "w", encoding="utf-8-sig", newline="") as f:
        f.write(f"# 학교명,구·군,개편 전 구·군,계열,설립,남녀   (출처: {os.path.basename(a.xlsx)} 고등학교 시트 · 구·군 기준: {a.district_col})\n")
        f.write("# fetch_ice_curriculum.py 는 앞의 두 칸(학교명,구·군)만 읽습니다. sync_school_list.py 로 다시 만들 수 있어요.\n")
        w = csv.writer(f, lineterminator="\n")
        for o in sorted(official, key=lambda o: (o[key], o["name"])):
            w.writerow([o["name"], o[key], o["gu"], o["cat"], o["est"], o["sex"]])
        for alias, base in ALIASES.items():
            if base in by_name:
                b = by_name[base]; w.writerow([alias, b[key], b["gu"], b["cat"], b["est"], b["sex"]])

    # 2) schools.json
    j = {"version": "", "operatorUrl": "", "note": "", "regions": []}
    if os.path.exists(a.json):
        with open(a.json, encoding="utf-8") as f:
            j = json.load(f)
    existing = {s["name"]: s for r in j.get("regions", []) for d in r["districts"] for s in d["schools"]}
    used = {s["id"] for s in existing.values()}

    def district_of(name):
        o = by_name.get(name) or by_name.get(ALIASES.get(name, ""))
        return o[key] if o else None

    targets = [o["name"] for o in official if o["cat"] in include]
    targets += [n for n in existing if n not in by_name]      # 공식 현황 밖 항목(국제계열 등)도 유지
    dists, moved, added, unknown = {}, [], [], []
    for name in targets:
        s = dict(existing.get(name) or {})
        d = district_of(name)
        if not d:
            unknown.append(name); d = "미분류"
        if not s:
            s = {"id": make_id(name, used), "name": name, "short": short_name(name), "file": "", "updated": ""}
            added.append(name)
        o = by_name.get(name) or by_name.get(ALIASES.get(name, ""))
        if o: s["cat"] = o["cat"]
        # 엑셀이 실제로 있는지로 선택 가능 여부 결정
        f = s.get("file") or f"{a.schools_dir}/{s['id']}.xlsx"
        if os.path.exists(f):
            s["file"] = f; s.pop("pending", None)
        else:
            s["file"] = ""; s["pending"] = True
        old = next((dn for r in j.get("regions", []) for dd in r["districts"] for x in dd["schools"] if x["name"] == name for dn in [dd["name"]]), None)
        if old and old != d:
            moved.append((name, old, d))
        dists.setdefault(d, []).append(s)

    region = {"name": "인천", "districts": [{"name": dn, "schools": sorted(v, key=lambda s: s["name"])}
                                           for dn, v in sorted(dists.items(), key=lambda kv: (kv[0] == "미분류", kv[0]))]}
    j["regions"] = [region]
    j["version"] = time.strftime("%Y-%m-%d") + "-official"
    j["note"] = (f"학교·구·군은 교육청 학교 현황({os.path.basename(a.xlsx)}) 기준, 구·군은 {a.district_col}. "
                 "편성표 엑셀이 없는 학교는 pending(편성표 준비 중)으로 표시됩니다. "
                 "학교를 추가하려면 schools/ 폴더에 엑셀을 올리고 tools/sync_school_list.py 를 다시 실행하세요.")
    with open(a.json, "w", encoding="utf-8") as f:
        json.dump(j, f, ensure_ascii=False, indent=2)

    # 3) 편성표 없는 학교
    with open(a.missing, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, lineterminator="\n"); w.writerow(["학교명", "id", "구·군", "계열", "사이트 표시"])
        for dd in region["districts"]:
            for s in dd["schools"]:
                if s.get("pending"):
                    w.writerow([s["name"], s["id"], dd["name"], s.get("cat", ""), "편성표 준비 중"])
        for o in sorted(official, key=lambda o: (o[key], o["name"])):
            if o["cat"] not in include:
                w.writerow([o["name"], "", o[key], o["cat"], "목록 제외(--include 로 추가 가능)"])

    total = sum(len(dd["schools"]) for dd in region["districts"])
    ready = sum(1 for dd in region["districts"] for s in dd["schools"] if not s.get("pending"))
    print(f"▶ 공식 현황 {len(official)}개교 읽음 · 구·군 기준: {a.district_col}")
    print(f"▶ schools.json: {total}개교 ({len(region['districts'])}개 구·군) · 선택 가능 {ready} · 편성표 준비 중 {total - ready}")
    for dd in region["districts"]:
        print(f"   {dd['name']:<6} {len(dd['schools']):>3}교  (준비 중 {sum(1 for s in dd['schools'] if s.get('pending'))})")
    if added: print(f"▶ 새로 넣은 학교 {len(added)}: " + ", ".join(added))
    if moved: print(f"▶ 구·군 바로잡음 {len(moved)}: " + ", ".join(f"{n}({o}→{d})" for n, o, d in moved))
    if unknown: print(f"▶ 공식 현황에 없는 학교(미분류) {len(unknown)}: " + ", ".join(unknown))
    print(f"▶ 기록: {a.districts}, {a.missing}")


if __name__ == "__main__":
    main()
